#!/usr/bin/env python3
"""
4E.3: Add RACI Matrix tab to governance-model.xlsx
Also adds Instructions tab per GOVERNANCE-TEMPLATES-AUDIT-SPEC.md
"""
import os
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FILE = os.path.join(BASE, "Templates", "04-planning", "governance-model.xlsx")

RED   = "E31C3D"
NAVY  = "1F3864"
WHITE = "FFFFFF"
LGRAY = "F2F2F2"
DGRAY = "D9D9D9"
AMBER = "FFC000"
GREEN_BG = "C6EFCE"
GREEN_FG = "006100"

def _fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def _font(bold=False, color=WHITE, size=10, italic=False):
    return Font(bold=bold, color=color, size=size, italic=italic, name="Calibri")

def _align(h="left", v="center", wrap=True):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def _border():
    s = Side(style="thin", color="AAAAAA")
    return Border(left=s, right=s, top=s, bottom=s)

def _hdr(ws, row, col, text, bg=RED, fg=WHITE, bold=True, size=10):
    c = ws.cell(row=row, column=col, value=text)
    c.fill = _fill(bg)
    c.font = _font(bold=bold, color=fg, size=size)
    c.alignment = _align()
    c.border = _border()
    return c

def _cell(ws, row, col, text, bg=WHITE, fg="000000", bold=False, wrap=True, align="left"):
    c = ws.cell(row=row, column=col, value=text)
    c.fill = _fill(bg)
    c.font = _font(bold=bold, color=fg)
    c.alignment = _align(h=align, wrap=wrap)
    c.border = _border()
    return c

def _title(ws, row, text, cols, bg=NAVY):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    c = ws.cell(row=row, column=1)
    c.value = text
    c.fill = _fill(bg)
    c.font = _font(bold=True, size=11)
    c.alignment = _align(h="center")

def set_col_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


ROLES = [
    "Engagement Lead Arch",
    "Platform Arch (Azure)",
    "Platform Arch (AWS/GCP)",
    "Delivery Mgr / PM",
    "Pre-Sales Architect",
    "Alliance Manager",
    "Oracle Practice Lead",
    "Customer IT Director",
    "App Owners",
    "Customer PM",
    "MS/AWS/GCP SE",
]

# Colour per RACI letter
RACI_BG = {
    "A/R": "C6EFCE",
    "R":   "BDD7EE",
    "A":   "FFEB9C",
    "C":   "E2EFDA",
    "I":   "F2F2F2",
    "N/A": "D9D9D9",
    "":    "FFFFFF",
}
RACI_FG = {
    "A/R": "006100",
    "R":   "1F497D",
    "A":   "9C5700",
    "C":   "375623",
    "I":   "595959",
    "N/A": "595959",
    "":    "FFFFFF",
}

def _raci(ws, row, col, value):
    bg = RACI_BG.get(value, WHITE)
    fg = RACI_FG.get(value, "000000")
    c = ws.cell(row=row, column=col, value=value)
    c.fill = _fill(bg)
    c.font = _font(bold=(value in ("A/R", "R", "A")), color=fg, size=9)
    c.alignment = _align(h="center", wrap=False)
    c.border = _border()


def add_instructions(wb):
    if "Instructions" in wb.sheetnames:
        del wb["Instructions"]
    ws = wb.create_sheet("Instructions", 0)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 80

    _title(ws, 1, "GOVERNANCE MODEL & RACI MATRIX", 2)
    _title(ws, 2, "CRA Framework -- Phase 1: Kickoff / Phase 4: Planning", 2, bg="404040")

    rows = [
        ("PURPOSE", ""),
        ("", "Define roles, responsibilities, and accountabilities for all CRA engagement activities."),
        ("", "The RACI matrix prevents ownership gaps and 'we did not know we owned that' disputes"),
        ("", "mid-engagement. Share and agree with the customer at the Phase 1 kickoff meeting."),
        ("", ""),
        ("WHO FILLS THIS IN", ""),
        ("", "Delivery Manager (RACI framework) + Engagement Lead Architect (technical rows)"),
        ("", "Reviewed and agreed with: Customer IT Director and Customer PM"),
        ("", ""),
        ("WHEN", ""),
        ("", "Phase 1 kickoff meeting. Update at each phase boundary if roles change."),
        ("", ""),
        ("STEP-BY-STEP", ""),
        ("", "1. Open the RACI Matrix tab"),
        ("", "2. Add actual names to role columns (replace [Lead Arch] with the named architect)"),
        ("", "3. Review each row with the customer -- confirm their A (Accountable) assignments"),
        ("", "4. Flag any row where the customer does not have an identified owner -- these are risks"),
        ("", "5. Add engagement-specific rows for custom activities (e.g., Oracle escalation if in scope)"),
        ("", "6. Save and attach to the engagement kickoff pack"),
        ("", ""),
        ("RACI KEY", ""),
        ("R -- Responsible",   "Does the work. Can be multiple people."),
        ("A -- Accountable",   "Approves the output. ONE person maximum per row."),
        ("C -- Consulted",     "Provides input before completion."),
        ("I -- Informed",      "Notified after completion."),
        ("A/R",                "Both accountable and responsible (single owner doing the work)"),
        ("", ""),
        ("HOW THIS FEEDS DOWNSTREAM", ""),
        ("", "Phase 2-4: Use RACI to identify who must attend workshops or sign deliverables"),
        ("", "Risk Register: RACI gaps (rows with no A) become risks to log in risk-assessment.xlsx"),
        ("", "Part 2 SOW: RACI carries forward and expands for Part 2 delivery team"),
        ("", ""),
        ("QUESTIONS?", "Refer to docs/guides/00-architect-onboarding-guide.md"),
    ]

    for i, (label, val) in enumerate(rows, 3):
        if label:
            c = ws.cell(row=i, column=1, value=label)
            c.font = _font(bold=True, color="000000", size=10)
            c.fill = _fill(LGRAY)
            c.alignment = _align()
            c.border = _border()
        else:
            ws.cell(row=i, column=1).border = _border()
        c2 = ws.cell(row=i, column=2, value=val)
        c2.font = Font(name="Calibri", size=10, color="000000")
        c2.alignment = _align()
        c2.border = _border()

    ws.row_dimensions[1].height = 28


def add_raci_matrix(wb):
    if "RACI Matrix" in wb.sheetnames:
        del wb["RACI Matrix"]
    ws = wb.create_sheet("RACI Matrix")
    ws.sheet_view.showGridLines = False

    # Column widths: Activity col wide, role cols narrow
    ws.column_dimensions["A"].width = 40
    for i in range(2, 13):
        ws.column_dimensions[get_column_letter(i)].width = 11

    total_cols = 12

    # Legend
    _title(ws, 1, "RACI MATRIX -- CRA Engagement Roles & Responsibilities", total_cols)
    _title(ws, 2, "Agree with customer at Phase 1 kickoff. Update actual names in role column headers.", total_cols, bg="404040")

    # RACI legend row
    legend_items = [("R = Responsible", "BDD7EE"), ("A = Accountable", "FFEB9C"),
                    ("C = Consulted", "E2EFDA"), ("I = Informed", "F2F2F2"),
                    ("A/R = Accountable+Responsible", "C6EFCE")]
    r = 3
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    _cell(ws, r, 1, "RACI KEY:", bg=DGRAY, bold=True, fg="000000")
    for i, (lbl, bg) in enumerate(legend_items):
        c = 3 + i
        _cell(ws, r, c, lbl, bg=bg, fg="000000", align="center")
    for c in range(8, total_cols + 1):
        _cell(ws, r, c, "", bg=DGRAY)

    # ── PHASE 1 ──
    r = 5
    _title(ws, r, "PHASE 1: Discovery", total_cols, bg="C00000")

    r = 6
    _cell(ws, r, 1, "Activity", bg=NAVY, bold=True, fg=WHITE)
    for ci, role in enumerate(ROLES[:10], 2):
        ws.cell(row=r, column=ci).value = f"[{role}]"
        ws.cell(row=r, column=ci).fill = _fill(NAVY)
        ws.cell(row=r, column=ci).font = _font(bold=True, size=8)
        ws.cell(row=r, column=ci).alignment = _align(h="center")
        ws.cell(row=r, column=ci).border = _border()
    ws.cell(row=r, column=12).value = "[MS/AWS/GCP SE]"
    ws.cell(row=r, column=12).fill = _fill(NAVY)
    ws.cell(row=r, column=12).font = _font(bold=True, size=8)
    ws.cell(row=r, column=12).alignment = _align(h="center")
    ws.cell(row=r, column=12).border = _border()

    # Roles columns: Lead, PlatAz, PlatAWSGCP, DM, PreSales, Alliance, Oracle, CustIT, AppOwn, CustPM, SE
    phase1_activities = [
        # (Activity, Lead, PlatAz, PlatAWSGCP, DM, PreSales, Alliance, Oracle, CustIT, AppOwn, CustPM, SE)
        ("Phase 1 Kickoff Meeting",             "A/R", "C", "C", "C", "C", "I", "I", "C", "I", "I", "I"),
        ("Application Scoping (initial pass)",  "R",   "C", "I", "I", "I", "I", "I", "A", "C", "I", "I"),
        ("Application Owner Validation",        "C",   "I", "I", "C", "I", "I", "I", "A", "R", "C", "I"),
        ("Discovery Tooling Deployment",        "R",   "R", "C", "I", "I", "I", "I", "C", "I", "I", "I"),
        ("CAB Firewall Change Request",         "C",   "R", "I", "I", "I", "I", "I", "A", "I", "R", "I"),
        ("Infrastructure Profiling",            "R",   "R", "C", "I", "I", "I", "I", "C", "C", "I", "I"),
        ("Dependency Mapping",                  "R",   "C", "I", "I", "I", "I", "I", "C", "R", "I", "I"),
        ("Oracle Flag -- Practice Escalation",  "A",   "R", "I", "I", "I", "I", "C", "I", "I", "I", "I"),
        ("Alliance Partner Pre-Registration",   "I",   "I", "I", "I", "C", "R", "I", "I", "I", "I", "I"),
        ("Phase 1 Deliverables Review",         "A",   "C", "C", "C", "I", "I", "I", "R", "I", "I", "I"),
        ("Phase 1 Exit Gate Sign-Off",          "A",   "C", "C", "C", "I", "I", "I", "R", "I", "I", "I"),
    ]
    bg_cycle = [WHITE, LGRAY]
    for i, row_data in enumerate(phase1_activities):
        r2 = 7 + i
        bg = bg_cycle[i % 2]
        _cell(ws, r2, 1, row_data[0], bg=bg, fg="000000", bold=True)
        for ci, val in enumerate(row_data[1:], 2):
            _raci(ws, r2, ci, val)

    # ── PHASE 2 ──
    r = 19
    _title(ws, r, "PHASE 2: Analysis", total_cols, bg="C00000")
    r = 20
    # Reuse same header
    _cell(ws, r, 1, "Activity", bg=NAVY, bold=True, fg=WHITE)
    for ci, role in enumerate(ROLES[:10], 2):
        ws.cell(row=r, column=ci).value = f"[{role}]"
        ws.cell(row=r, column=ci).fill = _fill(NAVY)
        ws.cell(row=r, column=ci).font = _font(bold=True, size=8)
        ws.cell(row=r, column=ci).alignment = _align(h="center")
        ws.cell(row=r, column=ci).border = _border()
    ws.cell(row=r, column=12).value = "[MS/AWS/GCP SE]"
    ws.cell(row=r, column=12).fill = _fill(NAVY)
    ws.cell(row=r, column=12).font = _font(bold=True, size=8)
    ws.cell(row=r, column=12).alignment = _align(h="center")
    ws.cell(row=r, column=12).border = _border()

    phase2_activities = [
        ("Governance Workshop",              "R",   "C", "I", "I", "I", "I", "I", "A", "C", "C", "I"),
        ("Cloud Readiness Scoring",          "R",   "R", "C", "I", "I", "I", "C", "C", "C", "I", "I"),
        ("Compliance Assessment",            "R",   "C", "I", "I", "I", "I", "I", "A", "I", "I", "I"),
        ("Oracle Practice Analysis",         "C",   "I", "I", "I", "I", "I", "R", "C", "I", "I", "I"),
        ("OSS Licence Risk Review",          "R",   "C", "I", "I", "I", "I", "I", "C", "R", "I", "I"),
        ("Phase 2 Deliverables Review",      "A",   "C", "C", "C", "I", "I", "I", "R", "I", "I", "I"),
    ]
    for i, row_data in enumerate(phase2_activities):
        r2 = 21 + i
        bg = bg_cycle[i % 2]
        _cell(ws, r2, 1, row_data[0], bg=bg, fg="000000", bold=True)
        for ci, val in enumerate(row_data[1:], 2):
            _raci(ws, r2, ci, val)

    # ── PHASE 3 ──
    r = 28
    _title(ws, r, "PHASE 3: Hyperscaler Evaluation", total_cols, bg="C00000")
    r = 29
    _cell(ws, r, 1, "Activity", bg=NAVY, bold=True, fg=WHITE)
    for ci, role in enumerate(ROLES[:10], 2):
        ws.cell(row=r, column=ci).value = f"[{role}]"
        ws.cell(row=r, column=ci).fill = _fill(NAVY)
        ws.cell(row=r, column=ci).font = _font(bold=True, size=8)
        ws.cell(row=r, column=ci).alignment = _align(h="center")
        ws.cell(row=r, column=ci).border = _border()
    ws.cell(row=r, column=12).value = "[MS/AWS/GCP SE]"
    ws.cell(row=r, column=12).fill = _fill(NAVY)
    ws.cell(row=r, column=12).font = _font(bold=True, size=8)
    ws.cell(row=r, column=12).alignment = _align(h="center")
    ws.cell(row=r, column=12).border = _border()

    phase3_activities = [
        ("TCO Modelling (Azure)",               "R",   "R", "I", "I", "C", "I", "I", "I", "I", "I", "C"),
        ("TCO Modelling (AWS)",                 "R",   "I", "R", "I", "C", "I", "I", "I", "I", "I", "C"),
        ("TCO Modelling (GCP)",                 "R",   "I", "R", "I", "C", "I", "I", "I", "I", "I", "C"),
        ("Licensing Overlay",                   "R",   "C", "I", "I", "C", "I", "R", "C", "I", "I", "I"),
        ("AMM/MAP/PSO Credit Estimation",       "C",   "I", "I", "I", "C", "R", "I", "I", "I", "I", "C"),
        ("Hyperscaler Scoring Matrix",          "A/R", "C", "I", "I", "C", "I", "I", "C", "I", "I", "I"),
        ("Weighting Agreement with Customer",   "R",   "I", "I", "I", "I", "I", "I", "A", "I", "I", "I"),
        ("Phase 3 Playback Presentation",       "A/R", "C", "C", "C", "I", "C", "I", "R", "I", "I", "I"),
    ]
    for i, row_data in enumerate(phase3_activities):
        r2 = 30 + i
        bg = bg_cycle[i % 2]
        _cell(ws, r2, 1, row_data[0], bg=bg, fg="000000", bold=True)
        for ci, val in enumerate(row_data[1:], 2):
            _raci(ws, r2, ci, val)

    # ── PHASE 4 ──
    r = 39
    _title(ws, r, "PHASE 4: Planning & Part 2 Entry Point", total_cols, bg="C00000")
    r = 40
    _cell(ws, r, 1, "Activity", bg=NAVY, bold=True, fg=WHITE)
    for ci, role in enumerate(ROLES[:10], 2):
        ws.cell(row=r, column=ci).value = f"[{role}]"
        ws.cell(row=r, column=ci).fill = _fill(NAVY)
        ws.cell(row=r, column=ci).font = _font(bold=True, size=8)
        ws.cell(row=r, column=ci).alignment = _align(h="center")
        ws.cell(row=r, column=ci).border = _border()
    ws.cell(row=r, column=12).value = "[MS/AWS/GCP SE]"
    ws.cell(row=r, column=12).fill = _fill(NAVY)
    ws.cell(row=r, column=12).font = _font(bold=True, size=8)
    ws.cell(row=r, column=12).alignment = _align(h="center")
    ws.cell(row=r, column=12).border = _border()

    phase4_activities = [
        ("Wave Plan Development",               "R",   "C", "I", "I", "I", "I", "I", "C", "C", "I", "I"),
        ("Risk Register (final)",               "R",   "C", "I", "I", "I", "I", "I", "C", "I", "I", "I"),
        ("Part 2 Entry Point Document",         "R",   "C", "I", "I", "R", "R", "C", "C", "I", "I", "I"),
        ("AMM/MAP/PSO Deal Registration",       "C",   "I", "I", "I", "I", "A/R","I", "I", "I", "I", "I"),
        ("Part 2 SOW Development",              "C",   "I", "I", "I", "A/R","C", "I", "C", "I", "I", "I"),
        ("Part 2 SOW -- Customer Signature",    "I",   "I", "I", "I", "I", "I", "I", "A", "I", "I", "I"),
    ]
    for i, row_data in enumerate(phase4_activities):
        r2 = 41 + i
        bg = bg_cycle[i % 2]
        _cell(ws, r2, 1, row_data[0], bg=bg, fg="000000", bold=True)
        for ci, val in enumerate(row_data[1:], 2):
            _raci(ws, r2, ci, val)

    # Roles legend at bottom
    r = 49
    _title(ws, r, "ROLES REFERENCE", total_cols, bg=NAVY)
    role_descriptions = [
        ("Engagement Lead Architect",    "Rackspace", "Overall delivery accountability; customer-facing lead"),
        ("Platform Architect (Azure)",   "Rackspace", "Azure technical delivery"),
        ("Platform Architect (AWS/GCP)", "Rackspace", "Multi-cloud evaluation support"),
        ("Delivery Manager / PM",        "Rackspace", "Schedule, risk, stakeholder management"),
        ("Pre-Sales Architect",          "Rackspace", "Pricing, commercial, Part 2 scoping"),
        ("Alliance Manager",             "Rackspace", "AMM/MAP/PSO deal registration"),
        ("Oracle Practice Lead",         "Rackspace", "Oracle RAC/EE licensing analysis (when applicable)"),
        ("Customer IT Director",         "Customer",  "Customer-side technical decision-maker"),
        ("App Owners",                   "Customer",  "Application-specific data and validation"),
        ("Customer PM / Coordinator",    "Customer",  "Internal coordination, SME scheduling"),
        ("Microsoft / AWS / GCP SE",     "Partner",   "Hyperscaler technical validation (Phase 3)"),
    ]
    _hdr(ws, 50, 1, "Role")
    _hdr(ws, 50, 2, "Internal / External")
    ws.merge_cells(start_row=50, start_column=3, end_row=50, end_column=12)
    _hdr(ws, 50, 3, "Description")
    for i, (role, side, desc) in enumerate(role_descriptions):
        r2 = 51 + i
        bg = bg_cycle[i % 2]
        _cell(ws, r2, 1, role, bg=bg, bold=True, fg="000000")
        _cell(ws, r2, 2, side, bg=bg, fg="000000", align="center")
        ws.merge_cells(start_row=r2, start_column=3, end_row=r2, end_column=12)
        _cell(ws, r2, 3, desc, bg=bg, fg="444444")

    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "B7"


def main():
    wb = openpyxl.load_workbook(FILE)
    print(f"Loaded: {FILE}")
    print(f"Existing tabs: {wb.sheetnames}")

    add_instructions(wb)
    add_raci_matrix(wb)

    wb.save(FILE)
    print(f"Saved: {FILE}")
    print(f"Final tabs: {wb.sheetnames}")


if __name__ == "__main__":
    main()
