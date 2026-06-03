#!/usr/bin/env python3
"""
10.6(b-e): Add document header block to remaining CRA DOCX templates:
  - cra-phase1-report-template.docx
  - part2-entry-point-template.docx
  - sow-template.docx
  - governance-workshop-schedule.docx
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn
from lxml import etree

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

RAX_RED = RGBColor(0xE3, 0x1C, 0x3D)

# ── Per-template header block content (from TEMPLATE-DESIGN-SPEC.md 10.6) ───

TEMPLATES = [
    {
        "file": os.path.join(BASE, "Templates", "executive-reporting",
                             "cra-phase1-report-template.docx"),
        "label": "cra-phase1-report-template.docx",
        "doc_purpose": (
            "Phase 1 delivery report documenting the infrastructure discovery findings, "
            "application inventory, and data quality summary."
        ),
        "phase": "Phase 1: Discovery -- End-of-Phase Deliverable",
        "audience": "Customer IT Director + Rackspace Lead Architect",
        "when": "Phase 1 Week 6-7 (final week); presented at Phase 1 exit gate meeting",
        "time": "8-12 hours; 20-30 pages when complete",
        "version": "v2.0",
        "customise": [
            "Replace [Customer Name] and [Date] on cover page",
            "Update infrastructure summary table from infrastructure-profiling.xlsx Data Quality tab",
            "Update application inventory summary from application-scoping-profiling.xlsx Summary tab",
            "Fill in the Utilisation Data Quality section with actual collection period and coverage %",
            "Flag any Oracle or EoL OS findings in the Key Findings section",
        ],
        "important": None,
    },
    {
        "file": os.path.join(BASE, "Templates", "executive-reporting",
                             "part2-entry-point-template.docx"),
        "label": "part2-entry-point-template.docx",
        "doc_purpose": (
            "The commercial handoff document that initiates the Part 2 migration engagement. "
            "Contains: Part 2 scope, timeline, investment, and next steps. "
            "The customer signs this at the Phase 4 playback meeting."
        ),
        "phase": "Phase 4: Planning -- Part 2 Handoff",
        "audience": "Customer CTO/CFO (for signature) + Rackspace Pre-Sales (for pricing)",
        "when": "Phase 4 weeks 3-5; present at Phase 4 playback and leave for signature",
        "time": "6-8 hours (much of content comes from Phase 3 outputs); 8-12 pages",
        "version": "v1.0",
        "customise": [
            "Section 3 (Hyperscaler Recommendation): copy from Phase 3 report; do not re-derive",
            "Section 4 (TCO Summary): reference Phase 3 TCO model; use 3-year net figure",
            "Section 5 (Partner Funding): confirm AMM/MAP/PSO registration status with Alliance Manager",
            "Section 10 (Investment): obtain pricing from Pre-Sales; do not estimate without approval",
        ],
        "important": (
            "IMPORTANT: AMM/MAP/PSO deal registration must be COMPLETE before Rackspace "
            "countersigns this document. Check with the Alliance Manager before preparing Section 5. "
            "See docs/DMG-MEDIA-UK-LESSONS-LEARNED.md Lesson 5."
        ),
    },
    {
        "file": os.path.join(BASE, "Templates", "04-planning", "sow-template.docx"),
        "label": "sow-template.docx",
        "doc_purpose": (
            "Statement of Work for CRA Part 1 engagement. Signed by customer and Rackspace "
            "before Phase 1 begins. Defines scope, deliverables, fees, timeline."
        ),
        "phase": "Pre-Engagement -- Phase 1 Gate",
        "audience": "Customer commercial/legal team + Rackspace Pre-Sales and Legal",
        "when": "Before Phase 1 kickoff; minimum 5 business days before start date",
        "time": "3-5 hours (customisation); must be reviewed by Pre-Sales Director",
        "version": "v1.0",
        "customise": [
            "Section 2-3 (Scope): replace estimated VM/application count with customer-provided figures",
            "Section 5 (Scope Variance Clause): confirm the 15% threshold is in the final document",
            "Section 8 (Utilisation Gate): confirm minimum data collection period is specified",
            "Section 10 (Fees): Pre-Sales Director must approve all fee figures before circulation",
            "Section 12 (Alliance Partner): confirm AMM/MAP/PSO registration clause is present",
        ],
        "important": (
            "IMPORTANT: Three critical clauses must be present before this SOW is sent for signature: "
            "(1) Scope Variance clause -- protects Rackspace if CMDB data is wrong; "
            "(2) Utilisation Data Gate clause -- protects Phase 3 data quality; "
            "(3) Alliance Partner Registration clause -- ensures AMM/MAP/PSO is registered on time. "
            "See Templates/executive-reporting/REPORTING-TEMPLATES-AUDIT-SPEC.md Section 4D.5."
        ),
    },
    {
        "file": os.path.join(BASE, "Templates", "04-planning",
                             "governance-workshop-schedule.docx"),
        "label": "governance-workshop-schedule.docx",
        "doc_purpose": (
            "Workshop schedule for the cloud governance foundations alignment session "
            "with the customer's IT Director (Phase 2)."
        ),
        "phase": "Phase 2: Analysis",
        "audience": "Customer IT Director + Rackspace Lead Architect",
        "when": "Phase 2 Week 1; book the session in Phase 1 Week 6",
        "time": "1-2 hours to customise agenda; workshop itself is 2-3 hours",
        "version": "v1.0",
        "customise": [
            "Add customer name and contact names",
            "Adjust domain weighting based on customer sector "
            "(e.g., heavier Compliance for FCA/healthcare)",
            "Add any sector-specific governance domains (e.g., PCI for payments; HIPAA for healthcare)",
        ],
        "important": None,
    },
]


def _spacer(doc):
    return doc.add_paragraph("")._element


def build_header_block(doc, tpl):
    """Build the 10.6 header block table for a given template spec dict."""
    row_count = 10 if tpl["important"] else 9

    table = doc.add_table(rows=row_count, cols=2)
    try:
        table.style = "Light Shading"
    except Exception:
        pass

    # Header row
    hrow = table.rows[0]
    hcell = hrow.cells[0].merge(hrow.cells[1])
    hcell.text = "DOCUMENT INFORMATION"
    try:
        hcell.paragraphs[0].runs[0].bold = True
        hcell.paragraphs[0].runs[0].font.color.rgb = RAX_RED
    except Exception:
        pass

    data_rows = [
        ("Document Purpose:", tpl["doc_purpose"]),
        ("CRA Phase:",        tpl["phase"]),
        ("Target Audience:",  tpl["audience"]),
        ("When to Complete:", tpl["when"]),
        ("Est. Time:",        tpl["time"]),
        ("Version:",          tpl["version"]),
    ]
    for i, (label, value) in enumerate(data_rows):
        r = table.rows[i + 1]
        r.cells[0].text = label
        try:
            r.cells[0].paragraphs[0].runs[0].bold = True
        except Exception:
            pass
        r.cells[1].text = value

    # Customise row
    crow = table.rows[7]
    ccell = crow.cells[0].merge(crow.cells[1])
    customise_lines = ["WHAT TO CUSTOMISE"] + ["- " + c for c in tpl["customise"]]
    ccell.text = "\n".join(customise_lines)
    try:
        ccell.paragraphs[0].runs[0].bold = True
    except Exception:
        pass

    # Separator
    sep_row = table.rows[8]
    sep_cell = sep_row.cells[0].merge(sep_row.cells[1])
    sep_cell.text = (
        "Template version: {version} -- CRA Framework v2.0 -- "
        "Rackspace Cloud Solutions Architecture".format(version=tpl["version"])
    )

    # Important row (optional)
    if tpl["important"] and row_count == 10:
        irow = table.rows[9]
        icell = irow.cells[0].merge(irow.cells[1])
        icell.text = tpl["important"]
        try:
            icell.paragraphs[0].runs[0].bold = True
            icell.paragraphs[0].runs[0].font.color.rgb = RAX_RED
        except Exception:
            pass

    return table._element


def process_file(tpl):
    if not os.path.exists(tpl["file"]):
        print(f"  SKIP (not found): {tpl['label']}")
        return

    doc = Document(tpl["file"])
    body = doc.element.body
    texts = [(p.style.name, p.text.lower())
             for p in doc.paragraphs if p.text.strip()]

    already = any("document information" in t[1] for t in texts)
    if already:
        print(f"  SKIP (header block already present): {tpl['label']}")
        return

    print(f"  Adding header block: {tpl['label']}")
    header_elem = build_header_block(doc, tpl)
    spacer_elem = _spacer(doc)
    body.insert(0, spacer_elem)
    body.insert(0, header_elem)

    doc.save(tpl["file"])
    print(f"    Saved.")


def main():
    print("10.6: Adding header blocks to DOCX templates")
    print("=" * 60)
    for tpl in TEMPLATES:
        process_file(tpl)
    print("\nDone.")


if __name__ == "__main__":
    main()
